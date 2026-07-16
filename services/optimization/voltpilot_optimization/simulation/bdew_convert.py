"""Convert the operator's own BDEW profile publication into the simulation's
``SIM_BDEW_H25_JSON`` file.

License situation (captain 2026-07-16): the official BDEW publication
"Repräsentative Profile BDEW H25 G25 L25 P25 S25" is free to download from
bdew.de but carries NO clear redistribution grant (the site's Impressum limits
downloads to private/non-commercial use; redistribution needs written
consent). So the repo ships ONLY this converter - the operator downloads the
official xlsx themself and converts it locally; neither the workbook nor any
derived per-slot values are ever committed.

Publication structure (derived from the real file, verified 2026-07-16):

* One sheet per profile (``H25``/``G25``/``L25``/``P25``/``S25``) plus a
  ``Dynamisierung`` sheet. Each profile sheet holds a 12-month x 3-day-type
  matrix: a header row of month datetimes (one per column), below it the
  day-type row (``SA``/``FT``/``WT`` per month), below that 96 quarter-hour
  rows (``00:00-00:15`` ... ``23:45-00:00``). Values are kWh per quarter
  hour, normalized to a 1,000,000 kWh annual consumption ("normiert auf
  1 Mio kWh Jahresverbrauch").
* The ``Dynamisierung`` sheet (a text box + formula image, no cell data)
  states: H25, P25 and S25 MUST be multiplied with the dynamization function
  when expanding ("Bei Ausrollen der Profile ist die Dynamisierungsfunktion
  anzuwenden!"); G25 and L25 must NOT. The coefficients are identical to the
  classic VDEW H0 function::

      x = x0 * (-3.92E-10*t^4 + 3.20E-7*t^3 - 7.02E-5*t^2 + 2.10E-3*t + 1.24)

  with ``t`` = day of the year (1 on Jan 1 ... 365, or 366 in leap years).
  Following the classic VDEW application convention the factor is rounded to
  4 decimal places before multiplying.

Expansion rules (this module's substance):

* **Day type per date**: ``FT`` on Sundays and nationwide German public
  holidays (reusing :mod:`voltpilot_forecast.holidays`), ``SA`` on Saturdays
  and - per the classic VDEW application guide, carried over since the
  publication says "analog zum VDEW H0" - on Dec 24 and Dec 31 when they fall
  on a weekday, ``WT`` otherwise.
* **DST / slot indexing**: the tables are day-type templates in LOCAL time,
  and the shipped adapter (``profiles._bdew_h25_weights``) indexes the
  ``slots`` array by the WALL-CLOCK quarter hour of the Berlin year - its
  ``(local - year_start)`` subtraction is a same-tzinfo naive subtraction,
  i.e. ``slot_of_year = (day_of_year - 1) * 96 + quarter_of_day``. So the
  converter emits exactly ``days_in_year * 96`` entries in that order (which
  equals the real year's physical slot count - DST cancels within a year).
  The DST semantics then live entirely in the adapter's lookup: on the
  spring-forward day the entries for the non-existent local quarter hours
  02:00-02:45 are simply never read (92 physical slots that day); on the
  fall-back day the local quarter hours 02:00-02:45 occur twice and their
  entries are read twice (100 physical slots). The adapter rescales over the
  slots it actually reads, so the phantom/doubled entries never skew the
  annual energy.
* **Leap years**: Feb 29 uses the February columns like any other February
  day; the dynamization ``t`` runs to 366 as the publication defines.
* **Output**: ``{"slots": [w0, ...]}`` - one weight per wall-clock 15-min
  slot of the target Berlin calendar year. Weights are emitted in kW at the
  publication's 1-Mio-kWh normalization (kWh per quarter hour times 4), so
  the pre-scaling annual energy ``sum(w) * 0.25 h`` must land near
  1,000,000 kWh - validated here. The simulation adapter
  (:func:`voltpilot_optimization.simulation.profiles.load_series_kw`) rescales
  exactly to the requested annual kWh, so only proportionality matters
  downstream.
"""

from __future__ import annotations

import json
import math
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from voltpilot_forecast.holidays import is_german_holiday

PROFILE_IDS = ("H25", "G25", "L25", "P25", "S25")

# Per the publication's Dynamisierung sheet: H25/P25/S25 are dynamized,
# G25/L25 are explicitly NOT ("nicht mit der Dynamisierungsfunktion zu
# verrechnen").
DYNAMIZED_PROFILES = frozenset({"H25", "P25", "S25"})

DAY_TYPES = ("SA", "FT", "WT")

SLOTS_PER_DAY = 96
SLOT_HOURS = 0.25

# The publication normalizes each profile to 1 Mio kWh annual consumption; the
# expanded year must reproduce that within a small tolerance (the dynamization
# function's annual mean is close to but not exactly 1).
NORMALIZED_ANNUAL_KWH = 1_000_000.0
ANNUAL_TOLERANCE = 0.02

_TIME_LABEL = re.compile(r"^\d{2}:\d{2}-\d{2}:\d{2}$")


class BdewConvertError(ValueError):
    """The workbook does not match the expected publication structure."""


@dataclass(frozen=True)
class ProfileTable:
    """One profile sheet: 96 quarter-hour kWh values per (month, day type)."""

    profile: str
    values: dict[tuple[int, str], list[float]]  # (month 1..12, day type) -> 96

    def value(self, month: int, day_type: str, quarter: int) -> float:
        return self.values[(month, day_type)][quarter]


def read_profile_table(xlsx_path: str, profile: str) -> ProfileTable:
    """Parse one profile sheet of the official publication workbook."""
    if profile not in PROFILE_IDS:
        raise BdewConvertError(
            f"Unbekanntes Profil '{profile}' - verfügbar: {', '.join(PROFILE_IDS)}"
        )
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment-dependent
        raise BdewConvertError(
            "openpyxl fehlt - bitte `pip install 'voltpilot-optimization[convert]'`"
        ) from exc
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if profile not in workbook.sheetnames:
            raise BdewConvertError(
                f"Blatt '{profile}' fehlt in {xlsx_path} - "
                f"vorhanden: {', '.join(workbook.sheetnames)}"
            )
        rows = list(workbook[profile].iter_rows(values_only=True))
    finally:
        workbook.close()

    month_row_idx = _find_month_header_row(rows, xlsx_path)
    month_row = rows[month_row_idx]
    day_type_row = rows[month_row_idx + 1]

    columns: list[tuple[int, int, str]] = []  # (column index, month, day type)
    for j, day_type in enumerate(day_type_row):
        if not isinstance(day_type, str) or day_type.strip() not in DAY_TYPES:
            continue
        month_cell = month_row[j]
        if not isinstance(month_cell, datetime):
            raise BdewConvertError(
                f"Spalte {j + 1} hat den Tagestyp '{day_type}' aber keinen "
                f"Monats-Header darüber - unerwartetes Layout in {xlsx_path}"
            )
        columns.append((j, month_cell.month, day_type.strip()))
    expected = {(m, dt) for m in range(1, 13) for dt in DAY_TYPES}
    found = {(m, dt) for _, m, dt in columns}
    if found != expected:
        raise BdewConvertError(
            f"Blatt '{profile}' deckt nicht 12 Monate x SA/FT/WT ab "
            f"(gefunden: {len(found)} Kombinationen) - unerwartetes Layout"
        )

    data_rows = [
        row
        for row in rows[month_row_idx + 2 :]
        if len(row) > 1
        and isinstance(row[1], str)
        and _TIME_LABEL.match(row[1].strip())
    ]
    if len(data_rows) != SLOTS_PER_DAY:
        raise BdewConvertError(
            f"Blatt '{profile}' hat {len(data_rows)} Viertelstunden-Zeilen "
            f"statt {SLOTS_PER_DAY} - unerwartetes Layout"
        )

    values: dict[tuple[int, str], list[float]] = {
        (m, dt): [] for m in range(1, 13) for dt in DAY_TYPES
    }
    for quarter, row in enumerate(data_rows):
        for j, month, day_type in columns:
            cell = row[j] if j < len(row) else None
            if not isinstance(cell, (int, float)) or not math.isfinite(cell):
                raise BdewConvertError(
                    f"Blatt '{profile}', Viertelstunde {quarter + 1}, Monat "
                    f"{month}/{day_type}: kein Zahlenwert ({cell!r})"
                )
            if cell < 0:
                raise BdewConvertError(
                    f"Blatt '{profile}', Viertelstunde {quarter + 1}, Monat "
                    f"{month}/{day_type}: negativer Wert {cell}"
                )
            values[(month, day_type)].append(float(cell))
    return ProfileTable(profile=profile, values=values)


def _find_month_header_row(rows: list[tuple], xlsx_path: str) -> int:
    """The month header row = the first row carrying >= 12 datetime cells."""
    for i, row in enumerate(rows[:10]):
        if sum(1 for c in row if isinstance(c, datetime)) >= 12:
            return i
    raise BdewConvertError(
        f"Keine Monats-Header-Zeile in {xlsx_path} gefunden - ist das die "
        "offizielle BDEW-Veröffentlichung (Blätter H25/G25/L25/P25/S25)?"
    )


def day_type_for(day: date) -> str:
    """The publication's day type of a calendar date (see module docstring)."""
    if day.weekday() == 6 or is_german_holiday(day):
        return "FT"
    if day.weekday() == 5:
        return "SA"
    if day.month == 12 and day.day in (24, 31):
        # Classic VDEW application rule ("analog zum VDEW H0"): Christmas Eve
        # and New Year's Eve count as Saturdays unless they fall on a Sunday.
        return "SA"
    return "WT"


def dynamization_factor(day_of_year: int) -> float:
    """The VDEW-H0 dynamization factor for day-of-year ``t`` (1-based),
    rounded to 4 decimal places per the classic application convention."""
    t = float(day_of_year)
    return round(
        -3.92e-10 * t**4 + 3.20e-7 * t**3 - 7.02e-5 * t**2 + 2.10e-3 * t + 1.24,
        4,
    )


def year_days(year: int) -> list[date]:
    """Every calendar date of ``year`` in order."""
    first = date(year, 1, 1)
    return [first + timedelta(days=i) for i in range((date(year + 1, 1, 1) - first).days)]


def expand_year(table: ProfileTable, year: int) -> list[float]:
    """Expand the profile table onto the target year's wall-clock slot grid
    (see the module docstring's DST section): 96 entries per calendar date,
    indexed exactly like the adapter's wall-clock ``slot_of_year``.

    Returns kW weights at the publication's 1-Mio-kWh normalization (kWh per
    quarter hour x 4), dynamized when the profile requires it.
    """
    dynamized = table.profile in DYNAMIZED_PROFILES
    weights: list[float] = []
    for t, day in enumerate(year_days(year), start=1):
        day_values = table.values[(day.month, day_type_for(day))]
        factor = dynamization_factor(t) if dynamized else 1.0
        weights.extend(v * factor / SLOT_HOURS for v in day_values)
    return weights


@dataclass(frozen=True)
class ConversionReport:
    """The validation summary printed after a conversion."""

    profile: str
    year: int
    dynamized: bool
    slots: int
    min_weight: float
    max_weight: float
    mean_weight: float
    annual_kwh_prescale: float
    weekday_sunday_ratio: float
    day_type_days: dict[str, int]

    def lines(self) -> list[str]:
        return [
            f"profile {self.profile} -> year {self.year} "
            f"({'dynamized' if self.dynamized else 'static'})",
            f"slots: {self.slots}",
            f"weight kW (at 1 Mio kWh normalization): "
            f"min {self.min_weight:.3f} / mean {self.mean_weight:.3f} / "
            f"max {self.max_weight:.3f}",
            f"annual energy before rescaling: {self.annual_kwh_prescale:,.0f} kWh "
            f"(publication normalization: {NORMALIZED_ANNUAL_KWH:,.0f} kWh)",
            f"weekday/sunday mean-weight ratio: {self.weekday_sunday_ratio:.3f}",
            "day types: "
            + ", ".join(f"{dt} {n}" for dt, n in sorted(self.day_type_days.items())),
        ]


def convert(xlsx_path: str, profile: str, year: int) -> tuple[dict, ConversionReport]:
    """Full conversion: parse, expand, validate. Returns the JSON document for
    ``SIM_BDEW_H25_JSON`` plus the validation report."""
    table = read_profile_table(xlsx_path, profile)
    weights = expand_year(table, year)
    report = _validate(profile, year, weights)
    document = {
        "profile": profile,
        "year": year,
        "source": "converted locally from the official BDEW publication "
        "(not redistributable - do not commit)",
        "slots": weights,
    }
    return document, report


def _validate(profile: str, year: int, weights: list[float]) -> ConversionReport:
    expected_slots = len(year_days(year)) * SLOTS_PER_DAY
    if len(weights) != expected_slots:
        raise BdewConvertError(
            f"Slot-Anzahl {len(weights)} != erwartete {expected_slots} für {year}"
        )
    if any(not math.isfinite(w) or w < 0 for w in weights):
        raise BdewConvertError("Gewichte müssen endlich und nicht-negativ sein")
    annual_kwh = sum(weights) * SLOT_HOURS
    if abs(annual_kwh - NORMALIZED_ANNUAL_KWH) > ANNUAL_TOLERANCE * NORMALIZED_ANNUAL_KWH:
        raise BdewConvertError(
            f"Jahresenergie vor Skalierung {annual_kwh:,.0f} kWh weicht mehr als "
            f"{ANNUAL_TOLERANCE:.0%} von der Normierung "
            f"({NORMALIZED_ANNUAL_KWH:,.0f} kWh) ab - falsches Blatt/Format?"
        )

    weekday_sum = weekday_n = sunday_sum = sunday_n = 0.0
    day_type_days: dict[str, int] = {"SA": 0, "FT": 0, "WT": 0}
    for i, day in enumerate(year_days(year)):
        day_type = day_type_for(day)
        day_type_days[day_type] += 1
        day_sum = sum(weights[i * SLOTS_PER_DAY : (i + 1) * SLOTS_PER_DAY])
        if day_type == "WT":
            weekday_sum += day_sum
            weekday_n += SLOTS_PER_DAY
        elif day_type == "FT":
            sunday_sum += day_sum
            sunday_n += SLOTS_PER_DAY
    ratio = (weekday_sum / weekday_n) / (sunday_sum / sunday_n)
    return ConversionReport(
        profile=profile,
        year=year,
        dynamized=profile in DYNAMIZED_PROFILES,
        slots=len(weights),
        min_weight=min(weights),
        max_weight=max(weights),
        mean_weight=sum(weights) / len(weights),
        annual_kwh_prescale=annual_kwh,
        weekday_sunday_ratio=ratio,
        day_type_days=day_type_days,
    )


def run(xlsx_path: str, profile: str, year: int, out_path: str) -> int:
    """CLI body: convert and write the JSON (``-`` = stdout), report to stderr."""
    document, report = convert(xlsx_path, profile, year)
    payload = json.dumps(document)
    if out_path == "-":
        sys.stdout.write(payload + "\n")
    else:
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(payload + "\n")
    for line in report.lines():
        print(line, file=sys.stderr)
    if out_path != "-":
        print(
            f"written: {out_path} (point SIM_BDEW_H25_JSON at it; "
            "do not commit - the values derive from the BDEW publication)",
            file=sys.stderr,
        )
    return 0
